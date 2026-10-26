from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from typing import Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas import CloudLedgerRow

_THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_NUM = "#,##0.00"
_TITLE_FONT = Font(bold=True, size=14)
_HEADER_FONT = Font(bold=True)
_BOLD = Font(bold=True)


def filter_ledger_rows(
    rows: Sequence[CloudLedgerRow],
    *,
    customer: str,
    clouds: Optional[Iterable[str]] = None,
    period: Optional[str] = None,
) -> List[CloudLedgerRow]:
    cloud_set = {c for c in (clouds or []) if c and c != "全部"}
    out: List[CloudLedgerRow] = []
    for r in rows:
        if r.customer != customer:
            continue
        if cloud_set and r.cloud not in cloud_set:
            continue
        if period and period != "全部" and r.period != period:
            continue
        out.append(r)
    out.sort(key=lambda x: (x.period, x.cloud, x.account), reverse=True)
    return out


def _apply_border_row(ws: Worksheet, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        ws.cell(row=row, column=c).border = _THIN


def _write_tencent_sheet(ws: Worksheet, rows: Sequence[CloudLedgerRow]) -> None:
    """与客户结算「腾讯云」sheet 一致（截图格式）。"""
    ws["A1"] = "腾讯云"
    ws["A1"].font = _TITLE_FONT

    headers = [
        "业务所属期",
        "账号信息",
        "实际消耗（￥）",
        "服务费比例",
        "服务费金额（￥）",
        "结算总金额（￥）",
        "备注",
    ]
    ws.append(headers)  # row 2
    for cell in ws[2]:
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    _apply_border_row(ws, 2, 7)

    by_period: dict[str, List[CloudLedgerRow]] = defaultdict(list)
    for r in rows:
        by_period[r.period].append(r)
    periods = sorted(by_period.keys(), reverse=True)

    for pi, period in enumerate(periods):
        group = by_period[period]
        # 账号顺序：尽量主账号 / 程序化 / 网盟靠前，其余按名
        order = {"主账号": 0, "程序化账号": 1, "网盟账号": 2}
        group.sort(key=lambda x: (order.get(x.account, 99), x.account))

        sum_c = sum_f = sum_s = 0.0
        for r in group:
            c = float(r.consume or 0)
            f = float(r.fee_amount or 0)
            s = float(r.settle_total or 0)
            sum_c += c
            sum_f += f
            sum_s += s
            ratio = r.fee_ratio or "2%"
            remark = r.remark or "真实消耗已预充值"
            ws.append([r.period, r.account, c, ratio, f, s, remark])
            row_idx = ws.max_row
            _apply_border_row(ws, row_idx, 7)
            for col in (3, 5, 6):
                ws.cell(row=row_idx, column=col).number_format = _NUM

        ws.append(
            [
                "合计",
                None,
                round(sum_c, 2),
                "—",
                round(sum_f, 2),
                round(sum_s, 2),
                None,
            ]
        )
        row_idx = ws.max_row
        for cell in ws[row_idx]:
            cell.font = _BOLD
        _apply_border_row(ws, row_idx, 7)
        for col in (3, 5, 6):
            ws.cell(row=row_idx, column=col).number_format = _NUM

        if pi < len(periods) - 1:
            ws.append([])

    widths = [12, 14, 14, 12, 14, 14, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_no_fee_sheet(ws: Worksheet, cloud: str, rows: Sequence[CloudLedgerRow]) -> None:
    """阿里云 / 华为云 / Akamai 短表格式。"""
    unit = (rows[0].currency if rows else "CNY") or "CNY"
    unit_label = "USD" if unit == "USD" else "￥"
    ws["A1"] = f"{cloud}（{unit}）" if unit == "USD" else cloud
    ws["A1"].font = _TITLE_FONT
    headers = ["业务所属期", "账号信息", f"实际消耗（{unit_label}）", "币种", "备注"]
    ws.append(headers)
    for cell in ws[2]:
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    _apply_border_row(ws, 2, 5)

    by_period: dict[str, List[CloudLedgerRow]] = defaultdict(list)
    for r in rows:
        by_period[r.period].append(r)

    # 短表按账期倒序平铺，末尾一个总合计（与现有结算表一致）
    sum_c = 0.0
    for period in sorted(by_period.keys(), reverse=True):
        for r in sorted(by_period[period], key=lambda x: x.account):
            c = float(r.consume or 0)
            sum_c += c
            cur = r.currency or unit
            remark = r.remark or (
                "Linode 月结（USD）" if cur == "USD" else "真实消耗已预充值（无服务费）"
            )
            ws.append([r.period, r.account, c, cur, remark])
            row_idx = ws.max_row
            _apply_border_row(ws, row_idx, 5)
            ws.cell(row=row_idx, column=3).number_format = _NUM

    ws.append(["合计", None, round(sum_c, 2), unit, None])
    row_idx = ws.max_row
    for cell in ws[row_idx]:
        cell.font = _BOLD
    _apply_border_row(ws, row_idx, 5)
    ws.cell(row=row_idx, column=3).number_format = _NUM

    for i, w in enumerate([12, 28, 16, 8, 36], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _safe_sheet_title(name: str) -> str:
    bad = set(r"[]:*?/\\")
    t = "".join("_" if c in bad else c for c in name)[:31]
    return t or "Sheet"


def build_ledger_xlsx(
    rows: Sequence[CloudLedgerRow],
    *,
    customer: str,
    clouds: Sequence[str],
) -> bytes:
    del customer, clouds  # 按行内云厂商分 sheet，标题用云名
    wb = Workbook()
    # 去掉默认 sheet，按云重建
    default = wb.active
    wb.remove(default)

    by_cloud: dict[str, List[CloudLedgerRow]] = defaultdict(list)
    for r in rows:
        by_cloud[r.cloud].append(r)

    # 腾讯云优先，其余按名
    cloud_order = sorted(by_cloud.keys(), key=lambda c: (0 if c == "腾讯云" else 1, c))
    for cloud in cloud_order:
        group = by_cloud[cloud]
        ws = wb.create_sheet(_safe_sheet_title(cloud))
        if cloud == "腾讯云":
            _write_tencent_sheet(ws, group)
        else:
            _write_no_fee_sheet(ws, cloud, group)

    if not wb.sheetnames:
        ws = wb.create_sheet("空")
        ws["A1"] = "无数据"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
