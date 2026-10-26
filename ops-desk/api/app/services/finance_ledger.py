"""财务分类账：有实际来源的自动汇总，没有的保持空（null）。"""
from __future__ import annotations

import re
from collections import defaultdict
from typing import Any, Dict, List, Optional, Tuple

from app.config import settings
from app.db import connect
from app.services.customers import settlement_customer_to_id
from app.services.scan_settlement import scan_settlement_dir

# book: income_cny | income_usd | expense_cny | expense_usd | footnote

INCOME_CNY_PARTIES = [
    ("hehe", "西安远景网络科技有限公司"),
    ("yidian", "亿点宇宙（陕西）信息科技有限公司"),
    ("shanhe", "北京山河聚量科技有限公司"),
]

INCOME_USD_PARTIES = [
    ("far", "FAR VISION收入"),
    ("fengyee", "FENGYEE TECHNOLOGY (HK) LIMITED"),
]

EXPENSE_CNY_PARTIES = [
    ("aliyun", "阿里云计算有限公司"),
    ("tencent", "腾讯云计算（北京）有限责任公司"),
    ("huawei", "华为云计算技术有限公司"),
    ("shujubao", "贵州数据宝网络科技有限公司"),
]

EXPENSE_USD_PARTIES = [
    ("fulu_u", "垫付U给FULU"),
]

FOOTNOTES = [
    ("tencent_xinghan_cost", "腾讯云 · 陕西星瀚算力服务器成本"),
    ("aliyun_xinghan_cost", "阿里云 · 陕西星瀚算力服务器成本"),
]

_CUSTOMER_MAP = {
    "hehe": "hehe",
    "远景": "hehe",
    "西安远景": "hehe",
    "亿点": "yidian",
    "山河": "shanhe",
    "星瀚": "xinghan",
}

_CLOUD_PAY_MAP = {
    "aliyun": "aliyun",
    "阿里云": "aliyun",
    "tencent": "tencent",
    "腾讯云": "tencent",
    "huawei": "huawei",
    "华为云": "huawei",
    "shujubao": "shujubao",
    "数据宝": "shujubao",
}

# 结算表「云」列 → 支出分类账 party_key（仅 CNY 四家）
_CLOUD_SETTLE_MAP = {
    "阿里云": "aliyun",
    "腾讯云": "tencent",
    "华为云": "huawei",
    "数据宝": "shujubao",
    "贵州数据宝": "shujubao",
}


def _map_cloud_party(cloud: str) -> Optional[str]:
    c = (cloud or "").strip()
    if not c:
        return None
    if c in _CLOUD_SETTLE_MAP:
        return _CLOUD_SETTLE_MAP[c]
    if c in _CLOUD_PAY_MAP:
        return _CLOUD_PAY_MAP[c]
    for k, pk in _CLOUD_SETTLE_MAP.items():
        if k in c:
            return pk
    return None

_TRINITY_PARTY = {
    "far": "far",
    "fengyee": "fengyee",
}


def _month_to_period(month_key: str) -> str:
    return month_key.replace("-", "")[:6]


def _month_parts(month_key: str) -> Tuple[int, int]:
    y, m = month_key.split("-")
    return int(y), int(m)


def _num(v: Any) -> float:
    if v is None or v == "" or v == "-":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _nz(v: Optional[float]) -> float:
    return float(v or 0)


def _balance(
    opening: Optional[float],
    recharge: Optional[float],
    consume: Optional[float],
    fee: Optional[float] = None,
) -> Optional[float]:
    """都空则余额也空；有任一数字则按空当 0 计算。"""
    if opening is None and recharge is None and consume is None and fee is None:
        return None
    return round(_nz(opening) + _nz(recharge) - _nz(consume) - _nz(fee), 2)


def settlement_income_cny(month_key: str) -> Dict[str, Dict[str, float]]:
    """客户结算表 CNY → 收入消耗/服务费。"""
    period = _month_to_period(month_key)
    _files, rows = scan_settlement_dir(settings.resolved_settlement_dir())
    agg: Dict[str, Dict[str, float]] = defaultdict(lambda: {"consume": 0.0, "fee": 0.0})
    hit: Dict[str, bool] = defaultdict(bool)
    for row in rows:
        if str(row.period) != period:
            continue
        if (row.currency or "CNY").upper() != "CNY":
            continue
        cid = (settlement_customer_to_id(row.customer) or "").strip()
        if not cid:
            name = row.customer or ""
            for k, pk in _CUSTOMER_MAP.items():
                if k in name:
                    cid = pk
                    break
        if cid not in ("hehe", "yidian", "shanhe"):
            continue
        agg[cid]["consume"] += _num(row.consume)
        agg[cid]["fee"] += _num(row.fee_amount)
        hit[cid] = True
    return {
        k: {"consume": round(v["consume"], 2), "fee": round(v["fee"], 2)}
        for k, v in agg.items()
        if hit[k]
    }


def payments_expense_recharge(month_key: str) -> Dict[str, float]:
    """出付款按「用途年月」汇总 → 支出侧本月充值。"""
    year, month = _month_parts(month_key)
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT cloud, amount FROM payments
            WHERE period_year = ? AND period_month = ?
            """,
            (year, month),
        ).fetchall()
    out: Dict[str, float] = defaultdict(float)
    for r in rows:
        cloud = str(r["cloud"] or "")
        key = _map_cloud_party(cloud) or _CLOUD_PAY_MAP.get(cloud.lower())
        if not key:
            continue
        out[key] += _num(r["amount"])
    return {k: round(v, 2) for k, v in out.items()}


_CUSTOMER_DISPLAY = {
    "hehe": "远景（hehe）",
    "远景": "远景（hehe）",
    "亿点": "亿点",
    "山河（BID）": "山河",
    "山河": "山河",
    "星瀚": "星瀚",
}


def _customer_display(label: str) -> str:
    s = (label or "").strip() or "未知"
    if s in _CUSTOMER_DISPLAY:
        return _CUSTOMER_DISPLAY[s]
    for k, name in _CUSTOMER_DISPLAY.items():
        if k in s:
            return name
    return s


def settlement_expense_by_cloud(month_key: str) -> Dict[str, Dict[str, Any]]:
    """客户结算按云厂商汇总 → 支出消耗 + 各客户公司明细。"""
    period = _month_to_period(month_key)
    _files, rows = scan_settlement_dir(settings.resolved_settlement_dir())
    agg: Dict[str, float] = defaultdict(float)
    by_cust: Dict[str, Dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for row in rows:
        if str(row.period) != period:
            continue
        if (row.currency or "CNY").upper() != "CNY":
            continue
        pk = _map_cloud_party(row.cloud or "")
        if not pk:
            continue
        amt = _num(row.consume)
        agg[pk] += amt
        by_cust[pk][_customer_display(row.customer or "")] += amt
    out: Dict[str, Dict[str, Any]] = {}
    for pk, total in agg.items():
        detail = [
            {"name": name, "amount": round(v, 2)}
            for name, v in sorted(by_cust[pk].items(), key=lambda x: -x[1])
            if abs(v) > 1e-9
        ]
        out[pk] = {"consume": round(total, 2), "detail": detail}
    return out


def trinity_income_usd(month_key: str) -> Dict[str, float]:
    """Trinity Invoice 按 bill_to 汇总 USD 消耗。"""
    year, month = _month_parts(month_key)
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT bill_to_id, amount_usd, invoice_date, period_label
            FROM trinity_invoices
            """
        ).fetchall()
    out: Dict[str, float] = defaultdict(float)
    for r in rows:
        bid = str(r["bill_to_id"] or "").strip().lower()
        party = _TRINITY_PARTY.get(bid)
        if not party:
            continue
        label = str(r["period_label"] or "")
        inv_m = str(r["invoice_date"] or "")[:7]
        matched = False
        # period_label: 9-2026 / 09-2026 / 2026-09 / Sep 2026
        if re.search(rf"(?:^|[^\d])0?{month}[-/年]{year}", label) or re.search(
            rf"{year}[-/年]0?{month}", label
        ):
            matched = True
        elif inv_m == month_key:
            matched = True
        if not matched:
            continue
        out[party] += _num(r["amount_usd"])
    return {k: round(v, 2) for k, v in out.items()}


def xinghan_aliyun_consume(month_key: str) -> Optional[float]:
    period = _month_to_period(month_key)
    _files, rows = scan_settlement_dir(settings.resolved_settlement_dir())
    total = 0.0
    found = False
    for row in rows:
        if str(row.period) != period:
            continue
        cid = settlement_customer_to_id(row.customer) or ""
        if cid != "xinghan" and "星瀚" not in (row.customer or ""):
            continue
        if "阿里" not in (row.cloud or ""):
            continue
        total += _num(row.consume)
        found = True
    return round(total, 2) if found else None


def _load_overrides(month_key: str) -> Dict[Tuple[str, str], dict]:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT book, party_key, party_name, opening, recharge, consume, fee, remark, sort_order
            FROM finance_ledger_lines
            WHERE month_key = ?
            """,
            (month_key,),
        ).fetchall()
    return {(str(r["book"]), str(r["party_key"])): dict(r) for r in rows}


def _round_or_none(v: Optional[float]) -> Optional[float]:
    if v is None:
        return None
    return round(float(v), 2)


def _line(
    *,
    book: str,
    party_key: str,
    party_name: str,
    opening: Optional[float] = None,
    recharge: Optional[float] = None,
    consume: Optional[float] = None,
    fee: Optional[float] = None,
    remark: str = "",
    consume_source: str = "",
    fee_source: str = "",
    recharge_source: str = "",
    opening_source: str = "",
    consume_detail: Optional[List[dict]] = None,
    sort_order: int = 0,
    override: Optional[dict] = None,
) -> dict:
    o = override
    if o is not None:
        # 已保存覆盖：整行以库为准（空数字存 0 也算已确认）
        opening = _round_or_none(o.get("opening"))
        recharge = _round_or_none(o.get("recharge"))
        consume = _round_or_none(o.get("consume")) if o.get("consume") is not None else None
        fee = _round_or_none(o.get("fee")) if o.get("fee") is not None else None
        remark = str(o.get("remark") or "")
        if o.get("party_name"):
            party_name = str(o["party_name"])
        sort_order = int(o.get("sort_order") or sort_order)
        opening_source = recharge_source = consume_source = fee_source = "override"

    with_fee = book.startswith("income")
    return {
        "book": book,
        "party_key": party_key,
        "party_name": party_name,
        "opening": opening,
        "recharge": recharge,
        "consume": consume,
        "fee": fee if with_fee or fee is not None else fee,
        "balance": _balance(opening, recharge, consume, fee if with_fee else None),
        "remark": remark or "",
        "consume_source": consume_source,
        "fee_source": fee_source,
        "recharge_source": recharge_source,
        "opening_source": opening_source,
        "consume_detail": consume_detail or [],
        "sort_order": sort_order,
    }


def build_ledger(month_key: str) -> dict:
    overrides = _load_overrides(month_key)
    income_auto = settlement_income_cny(month_key)
    pay_auto = payments_expense_recharge(month_key)
    expense_consume = settlement_expense_by_cloud(month_key)
    usd_auto = trinity_income_usd(month_key)
    xh_ali = xinghan_aliyun_consume(month_key)

    income_cny: List[dict] = []
    for i, (pk, pname) in enumerate(INCOME_CNY_PARTIES):
        auto = income_auto.get(pk)
        income_cny.append(
            _line(
                book="income_cny",
                party_key=pk,
                party_name=pname,
                consume=auto["consume"] if auto else None,
                fee=auto["fee"] if auto else None,
                consume_source="settlement" if auto else "",
                fee_source="settlement" if auto else "",
                sort_order=i,
                override=overrides.get(("income_cny", pk)),
            )
        )

    income_usd: List[dict] = []
    for i, (pk, pname) in enumerate(INCOME_USD_PARTIES):
        amt = usd_auto.get(pk)
        income_usd.append(
            _line(
                book="income_usd",
                party_key=pk,
                party_name=pname,
                consume=amt,
                fee=None,
                consume_source="trinity" if amt is not None else "",
                sort_order=i,
                override=overrides.get(("income_usd", pk)),
            )
        )

    expense_cny: List[dict] = []
    for i, (pk, pname) in enumerate(EXPENSE_CNY_PARTIES):
        rc = pay_auto.get(pk)
        pack = expense_consume.get(pk) or {}
        cs = pack.get("consume") if pack else None
        detail = pack.get("detail") if pack else []
        expense_cny.append(
            _line(
                book="expense_cny",
                party_key=pk,
                party_name=pname,
                recharge=rc,
                consume=cs,
                recharge_source="payments" if rc is not None else "",
                consume_source="settlement" if cs is not None else "",
                consume_detail=detail,
                sort_order=i,
                override=overrides.get(("expense_cny", pk)),
            )
        )

    expense_usd: List[dict] = []
    for i, (pk, pname) in enumerate(EXPENSE_USD_PARTIES):
        expense_usd.append(
            _line(
                book="expense_usd",
                party_key=pk,
                party_name=pname,
                sort_order=i,
                override=overrides.get(("expense_usd", pk)),
            )
        )

    footnotes: List[dict] = []
    for i, (pk, pname) in enumerate(FOOTNOTES):
        c: Optional[float] = None
        src = ""
        if pk == "aliyun_xinghan_cost" and xh_ali is not None:
            c = xh_ali
            src = "settlement"
        footnotes.append(
            _line(
                book="footnote",
                party_key=pk,
                party_name=pname,
                consume=c,
                consume_source=src,
                sort_order=i,
                override=overrides.get(("footnote", pk)),
            )
        )

    def _tot(rows: List[dict], with_fee: bool) -> dict:
        def s(key: str) -> Optional[float]:
            vals = [r[key] for r in rows if r.get(key) is not None]
            if not vals and all(r.get(key) is None for r in rows):
                return None
            return round(sum(_nz(r.get(key)) for r in rows), 2)

        op, rc, cs = s("opening"), s("recharge"), s("consume")
        fe = s("fee") if with_fee else None
        return {
            "opening": op,
            "recharge": rc,
            "consume": cs,
            "fee": fe if with_fee else None,
            "balance": _balance(op, rc, cs, fe if with_fee else None),
        }

    return {
        "month_key": month_key,
        "sources": {
            "income_cny_consume": "客户结算表 · 按公司",
            "income_usd_consume": "Trinity 收款 Invoice",
            "expense_cny_consume": "客户结算表 · 按云厂商",
            "expense_cny_recharge": "出付款申请（用途月）",
            "footnote_aliyun_xinghan": "星瀚结算表·阿里云",
            "empty": "期初、收入充值、数据宝无行、腾讯星瀚成本、垫付U 等暂无来源",
        },
        "income_cny": income_cny,
        "income_cny_total": _tot(income_cny, True),
        "income_usd": income_usd,
        "income_usd_total": _tot(income_usd, True),
        "expense_cny": expense_cny,
        "expense_cny_total": _tot(expense_cny, False),
        "expense_usd": expense_usd,
        "expense_usd_total": _tot(expense_usd, False),
        "footnotes": footnotes,
    }


def upsert_line(
    month_key: str,
    *,
    book: str,
    party_key: str,
    party_name: str,
    opening: Optional[float],
    recharge: Optional[float],
    consume: Optional[float],
    fee: Optional[float],
    remark: str = "",
    sort_order: int = 0,
) -> None:
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO finance_ledger_lines (
              month_key, book, party_key, party_name,
              opening, recharge, consume, fee, remark, sort_order, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now', '+8 hours'))
            ON CONFLICT(month_key, book, party_key) DO UPDATE SET
              party_name = excluded.party_name,
              opening = excluded.opening,
              recharge = excluded.recharge,
              consume = excluded.consume,
              fee = excluded.fee,
              remark = excluded.remark,
              sort_order = excluded.sort_order,
              updated_at = datetime('now', '+8 hours')
            """,
            (
                month_key,
                book,
                party_key,
                party_name,
                opening,
                recharge,
                consume,
                fee,
                remark or "",
                sort_order,
            ),
        )


def export_xlsx(month_key: str) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    data = build_ledger(month_key)
    wb = Workbook()
    ws = wb.active
    ws.title = "财务分类账"
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    yellow = PatternFill("solid", fgColor="FFF2CC")

    def cell_val(v: Optional[float]):
        return "" if v is None else v

    def write_block(start_row: int, title: str, rows: List[dict], total: dict, with_fee: bool) -> int:
        ws.cell(start_row, 1, title).font = Font(bold=True, size=14)
        r = start_row + 1
        headers = ["序号", "公司名称", "期初余额", "本月充值"]
        headers += ["本月消耗", "服务费", "余额", "备注"] if with_fee else ["本月消耗", "余额", "备注"]
        for c, h in enumerate(headers, 1):
            cell = ws.cell(r, c, h)
            cell.fill = header_fill
            cell.font = Font(bold=True)
        r += 1
        for i, row in enumerate(rows, 1):
            ws.cell(r, 1, i)
            ws.cell(r, 2, row["party_name"])
            ws.cell(r, 3, cell_val(row["opening"]))
            ws.cell(r, 4, cell_val(row["recharge"]))
            if with_fee:
                ws.cell(r, 5, cell_val(row["consume"])).fill = yellow
                ws.cell(r, 6, cell_val(row["fee"])).fill = yellow
                ws.cell(r, 7, cell_val(row["balance"]))
                ws.cell(r, 8, row["remark"])
            else:
                ws.cell(r, 5, cell_val(row["consume"])).fill = yellow
                ws.cell(r, 6, cell_val(row["balance"]))
                ws.cell(r, 7, row["remark"])
            r += 1
        ws.cell(r, 1, "总计").font = Font(bold=True)
        ws.cell(r, 3, cell_val(total.get("opening")))
        ws.cell(r, 4, cell_val(total.get("recharge")))
        if with_fee:
            ws.cell(r, 5, cell_val(total.get("consume")))
            ws.cell(r, 6, cell_val(total.get("fee")))
            ws.cell(r, 7, cell_val(total.get("balance")))
        else:
            ws.cell(r, 5, cell_val(total.get("consume")))
            ws.cell(r, 6, cell_val(total.get("balance")))
        return r + 2

    row = 1
    ws.cell(row, 1, "收入分类账").font = Font(bold=True, size=16)
    row = 3
    row = write_block(row, f"收入 · 人民币 · {month_key}", data["income_cny"], data["income_cny_total"], True)
    row = write_block(row, f"收入 · 美元 · {month_key}", data["income_usd"], data["income_usd_total"], True)
    row = write_block(row, f"支出 · 人民币 · {month_key}", data["expense_cny"], data["expense_cny_total"], False)
    row = write_block(row, f"支出 · 美元 · {month_key}", data["expense_usd"], data["expense_usd_total"], False)
    ws.cell(row, 1, "附注").font = Font(bold=True)
    row += 1
    for fn in data["footnotes"]:
        ws.cell(row, 1, fn["party_name"])
        ws.cell(row, 2, cell_val(fn["consume"])).fill = yellow
        row += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
