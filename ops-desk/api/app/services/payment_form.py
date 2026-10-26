from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path
from typing import Any, Dict

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.services.cn_amount import rmb_upper


def _copy_sheet(src: Worksheet, title: str):
    from openpyxl import Workbook

    wb = Workbook()
    dst = wb.active
    assert dst is not None
    dst.title = title[:31]

    for row in src.iter_rows():
        for cell in row:
            d = dst.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                d.font = copy(cell.font)
                d.border = copy(cell.border)
                d.fill = copy(cell.fill)
                d.number_format = cell.number_format
                d.protection = copy(cell.protection)
                d.alignment = copy(cell.alignment)

    for merged in src.merged_cells.ranges:
        dst.merge_cells(str(merged))

    for col_letter, dim in src.column_dimensions.items():
        dst.column_dimensions[col_letter].width = dim.width

    for idx, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[idx].height = dim.height

    # 打印：A5 铺满一页（原模板 scale=77 在 A5 上偏小留白）
    dst.print_area = src.print_area
    dst_ps = dst.page_setup
    dst_ps.orientation = "portrait"
    dst_ps.paperSize = 11  # A5
    dst.sheet_properties.pageSetUpPr.fitToPage = True
    dst_ps.fitToWidth = 1
    dst_ps.fitToHeight = 1
    # 不用固定 77%；交给「整表一页」按 A5 放大铺满
    dst_ps.scale = None
    dst.print_options.horizontalCentered = True
    # 略收边距，减少四周空白
    dm = dst.page_margins
    dm.left, dm.right = 0.4, 0.25
    dm.top, dm.bottom = 0.35, 0.3
    dm.header, dm.footer = 0.2, 0.2
    return wb


def format_cn_date(d: date) -> str:
    return f" {d.year} 年  {d.month}月  {d.day}日"


def build_purpose(purpose_name: str, cloud_label: str, period_month: int) -> str:
    # 对齐历史：腾讯多用「X云 N月…」；阿里多用「X业务阿里云 N月…」
    if "阿里" in cloud_label:
        return f"{purpose_name}业务{cloud_label} {period_month}月费用资源费用充值"
    return f"{purpose_name}{cloud_label} {period_month}月费用资源费用充值"


def generate_payment_xlsx(
    template_path: Path,
    out_path: Path,
    *,
    payer_name: str,
    apply_date: date,
    department: str,
    applicant: str,
    payee_name: str,
    payee_account: str,
    payee_bank: str,
    amount: float,
    purpose: str,
    sheet_title: str,
) -> Dict[str, Any]:
    if not template_path.is_file():
        raise FileNotFoundError(f"template not found: {template_path}")

    src_wb = load_workbook(template_path)
    if "2.付款申请单" not in src_wb.sheetnames:
        raise ValueError("template missing sheet 2.付款申请单")
    src = src_wb["2.付款申请单"]
    wb = _copy_sheet(src, sheet_title or "付款申请单")
    ws = wb.active
    assert ws is not None

    ws["B2"] = payer_name
    ws["F2"] = format_cn_date(apply_date)
    ws["B3"] = department
    ws["F3"] = applicant
    ws["C4"] = payee_name
    ws["C5"] = payee_account
    ws["C6"] = payee_bank
    ws["C7"] = round(float(amount), 2)
    ws["C8"] = rmb_upper(float(amount))
    ws["B9"] = purpose

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    src_wb.close()

    return {
        "path": str(out_path),
        "amount_upper": rmb_upper(float(amount)),
        "apply_date_label": format_cn_date(apply_date),
        "purpose": purpose,
    }
