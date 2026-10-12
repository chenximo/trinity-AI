from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path
from typing import Any, Dict, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_PAYMENT_INSTRUCTIONS = (
    "Beneficiary Name: TRINITY AI INC\n"
    "Beneficiary Address：8 The Green. Ste A, in the Cityof Dover County of Kent Zip Code 19901\n"
    "Beneficiary's A/C No：2910760256\n"
    "Beneficiary Bank : JPMorgan Chase Bank\n"
    "Beneficiary Bank Address: 270 Park Ave, New York, NY 10017\n"
    "Swift Code: CHASUS33\n"
    "Routing number：021000021"
)


def _copy_sheet(src: Worksheet, title: str) -> Workbook:
    wb = Workbook()
    dst = wb.active
    assert dst is not None
    dst.title = title[:31]

    for row in src.iter_rows():
        for cell in row:
            d = dst.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                d.font = copy(cell.font)
                d.border = copy(cell.border)
                d.fill = copy(cell.fill)
                d.number_format = cell.number_format
                d.protection = copy(cell.protection)
                d.alignment = copy(cell.alignment)

    for merged in src.merged_cells.ranges:
        dst.merge_cells(str(merged))

    for col_letter, dim in src.column_dimensions.items():
        dst.column_dimensions[col_letter].width = dim.width

    for idx, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[idx].height = dim.height

    # 复制图片（章）若模板源 sheet 带图——openpyxl 对 drawing 支持有限，先跳过
    return wb


def pick_template_sheet(wb) -> Worksheet:
    """优先用 FAR 样例作版式源；否则第一个非章子 sheet。"""
    for name in ("FAR", "hehe9月账单", "hehe8月账单"):
        if name in wb.sheetnames:
            return wb[name]
    for name in wb.sheetnames:
        if name != "章子":
            return wb[name]
    raise ValueError("template has no usable sheet")


def next_invoice_no(conn, apply_date: date) -> str:
    prefix = apply_date.strftime("%Y%m%d")
    row = conn.execute(
        """
        SELECT invoice_no FROM trinity_invoices
        WHERE invoice_no LIKE ? || '%'
        ORDER BY invoice_no DESC LIMIT 1
        """,
        (prefix,),
    ).fetchone()
    if not row:
        return f"{prefix}001"
    last = str(row["invoice_no"])
    try:
        seq = int(last[len(prefix) :]) + 1
    except ValueError:
        seq = 1
    return f"{prefix}{seq:03d}"


def generate_trinity_invoice_xlsx(
    template_path: Path,
    out_path: Path,
    *,
    invoice_no: str,
    invoice_date: date,
    bill_to: str,
    address: str,
    description: str,
    campaign: str,
    period_label: str,
    amount_usd: float,
    sheet_title: str = "Invoice",
    payment_instructions: Optional[str] = None,
) -> Dict[str, Any]:
    if not template_path.is_file():
        raise FileNotFoundError(f"template not found: {template_path}")

    src_wb = load_workbook(template_path)
    src = pick_template_sheet(src_wb)
    wb = _copy_sheet(src, sheet_title or "Invoice")
    ws = wb.active
    assert ws is not None

    ws["A1"] = "Invoice"
    ws["C2"] = "Invoice No."
    ws["D2"] = invoice_no
    ws["C3"] = "Invoice Date"
    ws["D3"] = invoice_date.isoformat()
    ws["A4"] = f"Bill To:{bill_to}" if not bill_to.startswith("Bill To") else bill_to
    ws["A5"] = address if address.lower().startswith("address") else f"Address：{address}"
    ws["A6"] = "Description"
    ws["B6"] = "Campaign"
    ws["C6"] = "Period"
    ws["D6"] = "Amount/USD"
    ws["A7"] = description
    ws["B7"] = campaign or "-"
    ws["C7"] = period_label
    ws["D7"] = round(float(amount_usd), 2)
    ws["A8"] = "Total/USD"
    ws["D8"] = round(float(amount_usd), 2)
    ws["A11"] = "Payment Instructions："
    ws["A12"] = payment_instructions or DEFAULT_PAYMENT_INSTRUCTIONS

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    src_wb.close()

    return {
        "path": str(out_path),
        "invoice_no": invoice_no,
        "amount_usd": round(float(amount_usd), 2),
    }
