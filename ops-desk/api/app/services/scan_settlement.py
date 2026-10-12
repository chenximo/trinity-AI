from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from app.schemas import CloudLedgerRow


def _num(v: Any) -> Optional[float]:
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _str(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _customer_from_filename(name: str) -> str:
    base = name.replace(".xlsx", "").replace(".xlsm", "")
    for suffix in ("云业务消耗费用", "云业务消耗", "消耗费用"):
        if base.endswith(suffix):
            return base[: -len(suffix)] or base
    return base


def _row_key(customer: str, period: str, cloud: str, account: str) -> str:
    return f"{customer}|{period}|{cloud}|{account}"


def scan_settlement_dir(directory: Path) -> Tuple[List[str], List[CloudLedgerRow]]:
    if not directory.is_dir():
        return [], []

    files = sorted(
        p.name
        for p in directory.glob("*.xlsx")
        if not p.name.startswith("~$") and not p.name.startswith(".~")
    )
    rows: List[CloudLedgerRow] = []

    for name in files:
        path = directory / name
        customer = _customer_from_filename(name)
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception:
            continue
        if "明细" not in wb.sheetnames:
            wb.close()
            continue
        ws = wb["明细"]
        header_idx: Optional[int] = None
        headers: Dict[int, str] = {}
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = list(row)
            if not cells:
                continue
            joined = [str(c).strip() if c is not None else "" for c in cells]
            if "业务所属期" in joined and "云厂商" in joined:
                header_idx = i
                headers = {idx: joined[idx] for idx in range(len(joined))}
                continue
            if header_idx is None:
                continue
            period = _str(cells[0]) if len(cells) > 0 else None
            if not period or period in ("合计", "业务所属期"):
                continue
            cloud = _str(cells[1]) if len(cells) > 1 else None
            account = _str(cells[2]) if len(cells) > 2 else None
            if not cloud or not account:
                continue

            # Prefer header map when columns shift; fallback fixed order from README
            by_name: Dict[str, Any] = {}
            for idx, h in headers.items():
                if idx < len(cells):
                    by_name[h] = cells[idx]

            consume = _num(by_name.get("实际消耗（￥）", cells[3] if len(cells) > 3 else None))
            fee_ratio = _str(by_name.get("服务费比例", cells[4] if len(cells) > 4 else None))
            fee_amount = _num(by_name.get("服务费金额（￥）", cells[5] if len(cells) > 5 else None))
            settle_total = _num(by_name.get("结算总金额（￥）", cells[6] if len(cells) > 6 else None))
            remark = _str(by_name.get("备注", cells[7] if len(cells) > 7 else None))
            fee_paid_excel = _str(
                by_name.get("服务费是否支付", cells[8] if len(cells) > 8 else None)
            )

            rows.append(
                CloudLedgerRow(
                    row_key=_row_key(customer, period, cloud, account),
                    customer=customer,
                    source_file=name,
                    period=str(period),
                    cloud=cloud,
                    account=account,
                    consume=consume,
                    fee_ratio=fee_ratio,
                    fee_amount=fee_amount,
                    settle_total=settle_total,
                    remark=remark,
                    fee_paid_excel=fee_paid_excel,
                    fee_paid=fee_paid_excel,
                    fee_paid_source="excel" if fee_paid_excel else "empty",
                )
            )
        wb.close()

    # newest period first, then customer
    rows.sort(key=lambda r: (r.period, r.customer, r.cloud, r.account), reverse=True)
    return files, rows


def list_recon_index(directory: Path) -> List[Dict[str, Any]]:
    if not directory.is_dir():
        return []
    items: List[Dict[str, Any]] = []
    for child in sorted(directory.iterdir()):
        if not child.is_dir() or child.name.startswith("."):
            continue
        sub = sorted(
            p.name for p in child.iterdir() if not p.name.startswith(".") and p.name != ".DS_Store"
        )[:40]
        items.append(
            {
                "kind": child.name,
                "name": child.name,
                "path": str(child.relative_to(directory.parent.parent) if False else child),
                "children": sub,
            }
        )
    return items
