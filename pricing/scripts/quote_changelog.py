#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""商务长期表共用「更新」日志（与 L3b / L3a 对内同结构：日期/分类/说明/备注，新→旧）。"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEET_UPD = "更新"

_THIN = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)


def solid_fill(rgb6: str) -> PatternFill:
    h = str(rgb6).strip().lstrip("#").upper()
    if len(h) == 6:
        h = "FF" + h
    elif len(h) == 8:
        h = "FF" + h[2:]
    return PatternFill(fill_type="solid", fgColor=h, bgColor=h)


def fmt_dot_date(d: date | None = None) -> str:
    d = d or date.today()
    return f"{d.year}.{d.month}.{d.day}"


def _cell_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def fmt_id_list(ids: list[str], limit: int = 12) -> str:
    ids = list(ids)
    if len(ids) <= limit:
        return "、".join(ids)
    return "、".join(ids[:limit]) + f" 等共 {len(ids)} 款"


def load_changelog(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    try:
        real = path.resolve()
    except OSError:
        real = path
    if not real.exists():
        return []
    try:
        wb = openpyxl.load_workbook(real, data_only=True)
    except Exception:  # noqa: BLE001
        return []
    try:
        if SHEET_UPD not in wb.sheetnames:
            return []
        ws = wb[SHEET_UPD]
        rows: list[dict[str, str]] = []
        for i, row in enumerate(ws.iter_rows(max_col=4, values_only=True), 1):
            a, b, c, d = (list(row) + [None, None, None, None])[:4]
            if i <= 2:
                continue
            date_s = _cell_str(a)
            detail = _cell_str(c)
            if not date_s and not detail:
                continue
            if date_s == "日期":
                continue
            if "账单说明" in date_s or "本页记录" in date_s or "本表为" in date_s:
                continue
            rows.append(
                {
                    "date": date_s,
                    "category": _cell_str(b),
                    "detail": detail,
                    "note": _cell_str(d),
                }
            )
        return rows
    finally:
        wb.close()


def merge_changelog(
    history: list[dict[str, str]],
    new_entries: list[dict[str, str]],
) -> list[dict[str, str]]:
    seen = {(e["date"], e["category"], e["detail"]) for e in history}
    prepend = [
        e for e in new_entries if (e["date"], e["category"], e["detail"]) not in seen
    ]
    return prepend + history


def inventory_diff_entries(
    old_ids: set[str],
    new_ids: set[str],
    *,
    today: str,
    category: str,
    note_add: str = "重建入表",
    note_rm: str = "本轮未再纳入",
) -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    added = sorted(new_ids - old_ids)
    removed = sorted(old_ids - new_ids)
    if added:
        entries.append(
            {
                "date": today,
                "category": category,
                "detail": f"新增 {len(added)} 款：{fmt_id_list(added)}",
                "note": note_add,
            }
        )
    if removed:
        entries.append(
            {
                "date": today,
                "category": category,
                "detail": f"移出 {len(removed)} 款：{fmt_id_list(removed)}",
                "note": note_rm,
            }
        )
    return entries


def load_custom_quote_model_ids(
    path: Path,
    *,
    sheet: str = "仅有折扣",
    model_col: int = 3,
    data_start: int = 3,
) -> set[str]:
    """定制表「仅有折扣」模型 ID 集合（默认第 3 列）。"""
    if not path.exists():
        return set()
    try:
        wb = openpyxl.load_workbook(path.resolve(), data_only=True)
    except Exception:  # noqa: BLE001
        return set()
    try:
        if sheet not in wb.sheetnames:
            return set()
        ws = wb[sheet]
        out: set[str] = set()
        for r in range(data_start, ws.max_row + 1):
            mid = _cell_str(ws.cell(r, model_col).value)
            if mid:
                out.add(mid)
        return out
    finally:
        wb.close()


def write_changelog_sheet(
    wb,
    entries: list[dict[str, str]],
    *,
    banner: str,
    index: int = 0,
) -> None:
    if SHEET_UPD in wb.sheetnames:
        del wb[SHEET_UPD]
    ws = wb.create_sheet(SHEET_UPD, index)
    header_fill = solid_fill("1B4F72")
    hint_fill = solid_fill("FEF3C7")
    alt = solid_fill("F8FAFC")
    white_bold = Font(name="PingFang SC", size=11, bold=True, color="FFFFFF")
    body = Font(name="PingFang SC", size=10, color="1D2939")

    ws.merge_cells("A1:D1")
    ws["A1"] = banner
    ws["A1"].font = Font(name="PingFang SC", size=9, color="713F12")
    ws["A1"].fill = hint_fill
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 48

    for i, h in enumerate(("日期", "分类", "更新说明", "备注"), 1):
        cell = ws.cell(2, i, h)
        cell.font = white_bold
        cell.fill = header_fill
        cell.border = _THIN
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    for idx, ent in enumerate(entries):
        r = 3 + idx
        fill = alt if idx % 2 else None
        vals = (
            ent.get("date") or "",
            ent.get("category") or "",
            ent.get("detail") or "",
            ent.get("note") or "",
        )
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = body
            cell.border = _THIN
            cell.alignment = Alignment(
                horizontal="center" if c <= 2 else "left",
                vertical="center",
                wrap_text=True,
            )
            if fill is not None:
                cell.fill = fill
        ws.row_dimensions[r].height = 36 if len(vals[2]) > 60 else 22

    for i, w in enumerate([12, 10, 64, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


def ensure_changelog_on_workbook(
    path: Path,
    *,
    banner: str,
    seed: dict[str, str] | None = None,
) -> None:
    """给已有 xlsx 补「更新」页（已有则跳过）。"""
    if not path.exists():
        return
    history = load_changelog(path)
    if history:
        return
    wb = openpyxl.load_workbook(path)
    if SHEET_UPD in wb.sheetnames:
        wb.close()
        return
    entries = [seed] if seed else [
        {
            "date": fmt_dot_date(),
            "category": "总册",
            "detail": "启用「更新」累积日志（与商务总册同结构）",
            "note": "首启",
        }
    ]
    write_changelog_sheet(wb, entries, banner=banner, index=0)
    wb.save(path)
    wb.close()
