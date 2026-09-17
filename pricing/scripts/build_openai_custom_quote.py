#!/usr/bin/env python3
"""OpenAI 客户指定模型定制报价（gpt-5.6-sol / gpt-5.6-terra / gpt-5.6-luna / gpt-image-2）。

真源：`output/Trinity模型报价表.xlsx`（L3a · 对外·仅折扣）。
只抽取客户清单里的模型行，列结构与源表保持一致（含 ≥$5k / ≥$10k / ≥$50k 三档折扣），
输出到 `output/custom-quotes/`，不回写主流程。
"""
from __future__ import annotations

import sys
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import openpyxl
from openpyxl.utils import get_column_letter

from build_outward_quote_standard import style_cell  # noqa: E402
from custom_quote_workbook import (  # noqa: E402
    CUSTOM_QUOTES_DIR,
    DEFAULT_VENDOR_FILLS,
    DEFAULT_ROW_FILL,
    _write_discount_title_row,
    merge_runs,
    sheet_styles,
)
from xlsx_wechat_compat import patch_xlsx_for_wechat  # noqa: E402

ROOT = SCRIPTS.parent
SRC = ROOT / "output/Trinity模型报价表.xlsx"
OUT = CUSTOM_QUOTES_DIR / "Trinity模型报价表_OpenAI定制.xlsx"

# 客户清单（按客户给出的顺序）
TARGET_IDS = ["gpt-5.6-sol", "gpt-5.6-terra", "gpt-5.6-luna", "gpt-image-2"]

HEADERS = [
    "模态",
    "厂商",
    "模型 ID",
    "显示名",
    "刊例价·输入",
    "刊例价·输出",
    "刊例价·缓存",
    "对外·≥$5k",
    "对外·≥$10k",
    "对外·≥$50k",
]
DATA_START = 3


def load_source_rows() -> list[dict]:
    """按客户清单从对外报价表抽取行（模态/厂商为合并单元格，需向下填充）。"""
    wb = openpyxl.load_workbook(SRC, read_only=True)
    ws = wb["仅有折扣"]
    modal = vendor = None
    index: dict[str, list] = {}
    for row in ws.iter_rows(min_row=DATA_START, values_only=True):
        if row[0]:
            modal = row[0]
        if row[1]:
            vendor = row[1]
        mid = str(row[2] or "").strip()
        if not mid:
            continue
        index[mid] = [modal, vendor] + [
            row[i] if row[i] is not None else "—" for i in range(3, 10)
        ]

    rows: list[dict] = []
    missing: list[str] = []
    for mid in TARGET_IDS:
        if mid not in index:
            missing.append(mid)
            continue
        modal, vendor, display, inp, out, cache, t5, t10, t50 = index[mid]
        rows.append(
            {
                "modal": modal,
                "vendor": vendor,
                "model_id": mid,
                "display": display,
                "inp": inp,
                "out": out,
                "cache": cache,
                "t5": t5,
                "t10": t10,
                "t50": t50,
            }
        )
    if missing:
        raise SystemExit(f"源表 {SRC.name} 中找不到模型: {', '.join(missing)}")
    return rows


def write_workbook(rows: list[dict]) -> None:
    styles = sheet_styles()
    thin = styles["thin"]
    header_fill = styles["header_fill"]
    body_font = styles["body_font"]
    num_font = styles["num_font"]
    white_bold = styles["white_bold"]
    center = styles["center"]
    left = styles["left"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "仅有折扣"
    _write_discount_title_row(ws, row=1, n_cols=len(HEADERS))

    for i, h in enumerate(HEADERS, 1):
        style_cell(ws.cell(2, i, h), fill=header_fill, font=white_bold, align=center, border=thin)
    ws.row_dimensions[2].height = 24

    for idx, row in enumerate(rows):
        r = DATA_START + idx
        fill = DEFAULT_VENDOR_FILLS.get(row["vendor"], DEFAULT_ROW_FILL)
        vals = [
            row["modal"],
            row["vendor"],
            row["model_id"],
            row["display"],
            row["inp"],
            row["out"],
            row["cache"],
            row["t5"],
            row["t10"],
            row["t50"],
        ]
        for c, v in enumerate(vals, 1):
            style_cell(
                ws.cell(r, c, v),
                font=num_font if c >= 5 else body_font,
                align=center if c == 1 or c >= 5 else left,
                border=thin,
                fill=fill,
            )

    merge_runs(ws, 1, [r["modal"] for r in rows], DATA_START)
    merge_runs(ws, 2, [(r["modal"], r["vendor"]) for r in rows], DATA_START)

    for i, w in enumerate([8, 12, 26, 22, 14, 14, 14, 12, 12, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{DATA_START}"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    patch_xlsx_for_wechat(OUT)


def main() -> None:
    rows = load_source_rows()
    write_workbook(rows)
    print(f"wrote {OUT} rows={len(rows)}")
    for r in rows:
        print(f"  {r['modal']} | {r['vendor']} | {r['model_id']} | {r['t5']} / {r['t10']} / {r['t50']}")


if __name__ == "__main__":
    main()
