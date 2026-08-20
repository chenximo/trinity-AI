#!/usr/bin/env python3
"""Crosswalk Media-Hub 接入清单 ↔ Trinity 线上刊例。

表一「有折扣」：Trinity 已上架且对外·≥$50k 有折扣。
表二「原价或未有」：Trinity 已上架但原价，或尚未上架。
"""

from __future__ import annotations

import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "pricing" / "scripts"))

from build_outward_quote_standard import (  # noqa: E402
    TIERS,
    fmt_tier_zhe,
    load_commercial,
    public_tiers_for,
    resolve_family,
)

DOCX = ROOT / "pricing" / "media-hub-AI大模型接入清单.docx"
OUT_XLSX = ROOT / "pricing" / "output" / "media-hub-Trinity对照.xlsx"

# Media-Hub 编码 / 上游 ID → Trinity 模型 ID（须在 /v1/prices 线上刊例中）
TRINITY_BY_CODE: dict[str, str | None] = {
    "ark_seedance_2_0_mini": None,
    "ark_seedance_2_0": "seedance-2.0-overseas",
    "ark_seedance_2_0_fast": "seedance-2.0-fast",
    "ark_seedance_1_5_pro": None,
    "happyhorse_1_1": "happyhorse-1.1",
    "happyhorse": "happyhorse-1.0",
    "happyhorse_video_edit": None,
    "kling_v26": "kling-2.6",
    "kling_v16": None,
    "openai": "os-2",
    "bytedance": None,
    "luma_ray_2": None,
    "hailuo_2_3": "hailuo-2.3",
    "grok_imagine_video": None,
    "openrouter_veo_3_1": "veo-3.1",
    "openrouter_wan_2_7": None,
    "openrouter_wan_2_6": None,
    "doubao_image": "SI-5.0-lite",
    "doubao_text": None,
    "doubao_vision": None,
    "(无模型编码)": "gemini-2.5-pro",
}

TRINITY_BY_UPSTREAM: dict[str, str] = {
    "google/veo-3.1": "veo-3.1",
    "gemini-2.5-pro": "gemini-2.5-pro",
    "doubao-seedream-5-0-260128": "SI-5.0-lite",
    "doubao-seedream-5.0-lite": "SI-5.0-lite",
    "openai/sora-2-pro": "os-2",
    "minimax/hailuo-2.3": "hailuo-2.3",
    "kwaivgi/kling-v2.6": "kling-2.6",
    "dashscope-happyhorse-1.1": "happyhorse-1.1",
    "happyhorse": "happyhorse-1.0",
    "doubao-seedance-2-0-fast-250828": "seedance-2.0-fast",
    "doubao-seedance-2-0-250828": "seedance-2.0-overseas",
}

SECTION_MODALITY = {
    0: "文生视频",
    1: "图生视频",
    2: "视频编辑",
    3: "文生图",
    4: "文本/LLM",
    5: "TTS",
    6: "数字人",
}

ENTERPRISE_TIER_IDX = next(i for i, (_, label, _) in enumerate(TIERS) if label == "对外·≥$50k")
ENTERPRISE_LABEL = TIERS[ENTERPRISE_TIER_IDX][1]

# Media-Hub 对照表专项覆盖（≥$50k），优先于商务公开阶梯
DISCOUNT_OVERRIDE: dict[str, str] = {
    "happyhorse-1.1": "4折",
    "happyhorse-1.0": "6折",
}


@dataclass
class HubModel:
    code: str
    name: str
    upstream: str = ""
    supplier: str = ""
    modalities: set[str] = field(default_factory=set)


def _cell(row, idx: int) -> str:
    if idx >= len(row):
        return ""
    return str(row[idx] or "").strip()


def parse_docx(path: Path) -> dict[str, HubModel]:
    doc = Document(str(path))
    by_code: dict[str, HubModel] = {}

    for ti, table in enumerate(doc.tables):
        if ti > 6:
            break
        modality = SECTION_MODALITY.get(ti)
        if not modality:
            continue
        rows = table.rows
        if len(rows) < 2:
            continue
        header = [_cell(r.cells, i) for i, r in enumerate([rows[0]]) for r in [rows[0]]]
        # normalize header from first row cells
        header = [c.text.strip() for c in rows[0].cells]

        def col(name: str) -> int | None:
            for i, h in enumerate(header):
                if name in h:
                    return i
            return None

        idx_name = col("模型名称") or 1
        idx_code = col("模型编码")
        idx_upstream = col("上游模型 ID")
        idx_supplier = col("供应商")

        for row in rows[1:]:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            name = _cell(cells, idx_name)
            if not name or name == "模型名称":
                continue
            code = _cell(cells, idx_code) if idx_code is not None else ""
            if not code:
                code = f"__name__:{name}"
            upstream = _cell(cells, idx_upstream) if idx_upstream is not None else ""
            supplier = _cell(cells, idx_supplier) if idx_supplier is not None else ""

            if code not in by_code:
                by_code[code] = HubModel(
                    code=code,
                    name=name,
                    upstream=upstream,
                    supplier=supplier,
                )
            by_code[code].modalities.add(modality)
            if upstream and not by_code[code].upstream:
                by_code[code].upstream = upstream
            if supplier and not by_code[code].supplier:
                by_code[code].supplier = supplier
    return by_code


def resolve_trinity_id(hub: HubModel, listed: set[str]) -> str | None:
    # 显式标记为 None 的编码不再做模糊匹配
    if hub.code in TRINITY_BY_CODE:
        tid = TRINITY_BY_CODE[hub.code]
        if tid is None:
            return None
        if tid in listed:
            return tid
        return None

    if hub.upstream:
        tid = TRINITY_BY_UPSTREAM.get(hub.upstream)
        if tid and tid in listed:
            return tid

    return None


def modality_label(hub: HubModel) -> str:
    order = ["文生视频", "图生视频", "视频编辑", "文生图", "文本/LLM", "TTS", "数字人"]
    parts = [m for m in order if m in hub.modalities]
    return "、".join(parts) if parts else "—"


def load_listed_models() -> set[str]:
    import json

    listed: set[str] = set()
    for mod in ("text", "image", "video"):
        p = ROOT / "pricing" / "output" / "online" / f"prices-api-{mod}.json"
        raw = json.loads(p.read_text(encoding="utf-8"))
        for e in raw.get("data") or []:
            mid = str(e.get("model") or "").strip()
            if mid:
                listed.add(mid)
    return listed


def enterprise_discount(
    trinity_id: str,
    model_families: dict[str, list[float]],
    recommended: dict[str, float],
) -> str:
    if trinity_id in DISCOUNT_OVERRIDE:
        return DISCOUNT_OVERRIDE[trinity_id]
    fam = resolve_family(trinity_id, model_families, recommended)
    zhes = public_tiers_for(trinity_id, fam)
    if not zhes:
        return "—"
    zhe = zhes[ENTERPRISE_TIER_IDX]
    return fmt_tier_zhe(zhe)


def is_discounted(zhe_label: str) -> bool:
    """有折扣：非原价、非空、非 —。"""
    s = (zhe_label or "").strip()
    return bool(s) and s not in ("原价", "—", "-")


def write_workbook(discounted: list[dict], table2: list[dict], path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "有折扣"
    ws2 = wb.create_sheet("原价或未有")

    headers_have = [
        "多模态",
        "Media-Hub 模型",
        "Media-Hub 编码",
        "Trinity 模型 ID",
        ENTERPRISE_LABEL,
        "备注",
    ]
    headers_t2 = [
        "多模态",
        "Media-Hub 模型",
        "Media-Hub 编码",
        "状态",
        "Trinity 模型 ID",
        ENTERPRISE_LABEL,
        "上游模型 ID",
        "备注",
    ]

    for ws, headers, rows in (
        (ws1, headers_have, discounted),
        (ws2, headers_t2, table2),
    ):
        ws.append(headers)
        for r in rows:
            ws.append([r.get(h, "") for h in headers])
        widths = [18, 28, 22, 12, 22, 14, 28, 28][: len(headers)]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_hehe_sheet(wb, table2: list[dict], *, index: int | None = None) -> None:
    """写入「hehe公司需求」Sheet（Media-Hub 表二：原价或未有）。"""
    name = "hehe公司需求"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, index) if index is not None else wb.create_sheet(name)
    headers = [
        "多模态",
        "Media-Hub 模型",
        "Media-Hub 编码",
        "状态",
        "Trinity 模型 ID",
        ENTERPRISE_LABEL,
        "上游模型 ID",
        "备注",
    ]
    ws.append(headers)
    for r in table2:
        ws.append([r.get(h, "") for h in headers])
    widths = [18, 28, 22, 12, 22, 14, 28, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def collect_media_hub_rows() -> tuple[list[dict], list[dict]]:
    """返回 (有折扣, 原价或未有)。"""
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)

    hubs = parse_docx(DOCX)
    listed = load_listed_models()
    _, model_families, recommended = load_commercial()

    discounted: list[dict] = []
    table2: list[dict] = []

    notes = {
        "ark_seedance_2_0": "Trinity 对应海外线路 seedance-2.0-overseas",
        "openai": "Trinity 模型 ID：os-2（Sora 2）",
        "doubao_image": "Trinity 模型 ID：SI-5.0-lite",
        "(无模型编码)": "经 Trinity Desk Proxy 接入",
        "happyhorse": "对照表专项：≥$50k 调为 6折",
        "happyhorse_1_1": "对照表专项：≥$50k 调为 4折",
    }

    for hub in sorted(hubs.values(), key=lambda h: (modality_label(h), h.name)):
        tid = resolve_trinity_id(hub, listed)
        mod = modality_label(hub)
        code = hub.code if not hub.code.startswith("__name__:") else "—"
        if tid:
            zhe = enterprise_discount(tid, model_families, recommended)
            row_base = {
                "多模态": mod,
                "Media-Hub 模型": hub.name,
                "Media-Hub 编码": code,
                "Trinity 模型 ID": tid,
                ENTERPRISE_LABEL: zhe,
                "上游模型 ID": hub.upstream or "—",
                "备注": notes.get(hub.code, ""),
            }
            if is_discounted(zhe):
                discounted.append(row_base)
            else:
                table2.append(
                    {
                        **row_base,
                        "状态": "已有·原价",
                        "备注": notes.get(hub.code, "")
                        or "Trinity 已上架，对外·≥$50k 为原价",
                    }
                )
        else:
            miss_note = ""
            if hub.code == "ark_seedance_2_0_mini":
                miss_note = "Trinity 仅有 seedance-2.0-fast 系列，无 Mini 规格"
            elif hub.code == "kling_v16":
                miss_note = "Trinity 最低可灵为 kling-2.5-turbo / kling-2.6"
            elif hub.code.startswith("openrouter_wan"):
                miss_note = "Trinity 视频刊例暂无 Wan 系列"
            table2.append(
                {
                    "多模态": mod,
                    "Media-Hub 模型": hub.name,
                    "Media-Hub 编码": code,
                    "状态": "未有",
                    "Trinity 模型 ID": "—",
                    ENTERPRISE_LABEL: "—",
                    "上游模型 ID": hub.upstream or "—",
                    "备注": miss_note,
                }
            )

    table2.sort(
        key=lambda r: (
            0 if r.get("状态") == "已有·原价" else 1,
            r.get("多模态", ""),
            r.get("Media-Hub 模型", ""),
        )
    )
    return discounted, table2


def main() -> None:
    discounted, table2 = collect_media_hub_rows()
    write_workbook(discounted, table2, OUT_XLSX)
    print(f"有折扣: {len(discounted)}")
    print(f"原价或未有: {len(table2)}")
    print(f"written: {OUT_XLSX}")


if __name__ == "__main__":
    main()
